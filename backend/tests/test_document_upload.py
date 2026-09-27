"""Document upload + validation + version control (SRS Steps 4-8, vi, x)."""

import os
import tempfile

from tests.helpers import (
    SECURITY_POLICY_V1, SECURITY_POLICY_V2, bootstrap_admin, get_client, make_docx, make_pdf, upload, upload_version,
)


def _client():
    c = get_client()
    return c, bootstrap_admin(c), tempfile.mkdtemp()


def test_invalid_file_type_is_rejected():
    c, h, tmp = _client()
    p = os.path.join(tmp, "malware.exe")
    open(p, "wb").write(b"MZ binary")
    r = upload(c, h, p, "BAD-01")
    assert r.status_code == 422 and "Unsupported file type" in r.text


def test_empty_document_is_rejected():
    c, h, tmp = _client()
    p = os.path.join(tmp, "empty.txt")
    open(p, "w").write("")
    assert upload(c, h, p, "EMPTY-01").status_code == 422


def test_duplicate_content_is_rejected():
    c, h, tmp = _client()
    p = make_docx(os.path.join(tmp, "a.docx"), SECURITY_POLICY_V1)
    assert upload(c, h, p, "SEC-A").status_code == 200
    r = upload(c, h, p, "SEC-B")
    assert r.status_code == 409 and "SEC-A" in r.text


def test_duplicate_code_and_bad_metadata():
    c, h, tmp = _client()
    p1 = make_docx(os.path.join(tmp, "a.docx"), SECURITY_POLICY_V1)
    p2 = make_docx(os.path.join(tmp, "b.docx"), SECURITY_POLICY_V2)
    assert upload(c, h, p1, "SEC-A").status_code == 200
    assert upload(c, h, p2, "SEC-A").status_code == 422
    assert upload(c, h, p2, "SEC-C", doc_type="Rumour").status_code == 422
    assert upload(c, h, p2, "SEC-D", version="v-two").status_code == 422


def test_effective_and_expiry_dates_are_validated():
    c, h, tmp = _client()
    p = make_docx(os.path.join(tmp, "a.docx"), SECURITY_POLICY_V1)
    with open(p, "rb") as f:
        r = c.post("/documents/upload", headers=h, files={"file": ("a.docx", f)},
                   data={"document_code": "SEC-X", "title": "t", "doc_type": "Policy", "department": "All",
                         "version": "1.0", "effective_date": "2026-05-01", "expiry_date": "2026-01-01"})
    assert r.status_code == 422 and "expiry_date" in r.text


def test_pdf_docx_txt_and_md_are_parsed_with_traceability():
    c, h, tmp = _client()
    pdf = make_pdf(os.path.join(tmp, "p.pdf"), ["Leave Policy\nEmployees must apply for leave 3 days in advance.",
                                                "Section 2\nSick leave requires a medical certificate."])
    r = upload(c, h, pdf, "LEAVE-01", "Policy")
    assert r.status_code == 200, r.text
    chunks = c.get(f"/documents/{r.json()['document_id']}/chunks", headers=h).json()["chunks"]
    assert [ch["location"] for ch in chunks] == ["page 1", "page 2"]

    md = os.path.join(tmp, "faq.md")
    open(md, "w").write("# Refunds\nRefunds above PKR 10,000 are escalated.\n\n# Leave\nLeave needs approval.\n")
    r = upload(c, h, md, "FAQ-MD", "FAQ")
    assert r.status_code == 200, r.text
    chunks = c.get(f"/documents/{r.json()['document_id']}/chunks", headers=h).json()["chunks"]
    assert {ch["section"] for ch in chunks} == {"Refunds", "Leave"}

    import docx
    d = docx.Document()
    d.add_heading("Escalation matrix", level=1)
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Amount", "Escalate to"
    t.cell(1, 0).text, t.cell(1, 1).text = "Above PKR 10,000", "Team Leader"
    tp = os.path.join(tmp, "table.docx")
    d.save(tp)
    r = upload(c, h, tp, "SOP-TBL", "SOP")
    chunks = c.get(f"/documents/{r.json()['document_id']}/chunks", headers=h).json()["chunks"]
    assert any("Above PKR 10,000 | Team Leader" in ch["content"] for ch in chunks)


def test_new_version_supersedes_old_one():
    c, h, tmp = _client()
    v1 = make_docx(os.path.join(tmp, "v1.docx"), SECURITY_POLICY_V1)
    v2 = make_docx(os.path.join(tmp, "v2.docx"), SECURITY_POLICY_V2)
    assert upload(c, h, v1, "SEC-v1").status_code == 200
    r = upload_version(c, h, "SEC-v1", v2, "SEC-v2", "2.0")
    assert r.status_code == 200, r.text
    docs = {d["document_code"]: d for d in c.get("/documents/list", headers=h).json()}
    assert docs["SEC-v1"]["is_active_version"] is False and docs["SEC-v1"]["expiry_date"]
    assert docs["SEC-v2"]["is_active_version"] is True
    hist = c.get("/documents/versions/SEC-v1", headers=h).json()
    assert hist["total_versions"] == 2
