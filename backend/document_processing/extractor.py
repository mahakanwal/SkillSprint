"""
document_processing/extractor.py

Real, working extraction logic for uploaded PDF and DOCX files.
- PDF: extracted using pdfplumber (page-level text + page numbers retained)
- DOCX: extracted using python-docx (heading-aware, paragraph-level)

Both extractors return a common structure so the rest of the pipeline
(chunker.py) doesn't need to know which file type it came from.
"""

import os
import pdfplumber
from docx import Document as DocxDocument


class ExtractionError(Exception):
    """Raised when a document cannot be parsed."""
    pass


def extract_pdf(file_path: str) -> list[dict]:
    """
    Extracts text from a PDF file, page by page.

    Returns a list of dicts:
    [{"page_number": 1, "text": "..."}, {"page_number": 2, "text": "..."}, ...]
    """
    if not os.path.exists(file_path):
        raise ExtractionError(f"File not found: {file_path}")

    pages = []
    try:
        with pdfplumber.open(file_path) as pdf:
            if len(pdf.pages) == 0:
                raise ExtractionError("PDF has no pages (possibly corrupted or empty).")

            for i, page in enumerate(pdf.pages, start=1):
                text = page.extract_text() or ""
                pages.append({
                    "page_number": i,
                    "text": text.strip()
                })
    except Exception as e:
        raise ExtractionError(f"Failed to parse PDF '{file_path}': {str(e)}")

    non_empty_pages = [p for p in pages if p["text"]]
    if not non_empty_pages:
        raise ExtractionError(
            f"No extractable text found in '{file_path}'. "
            f"It may be a scanned/image-based PDF requiring OCR."
        )

    return pages


def extract_docx(file_path: str) -> list[dict]:
    """
    Extracts text from a DOCX file, tracking headings so later chunking
    can attach a 'Section/Heading' reference to each paragraph.

    Returns a list of dicts:
    [{"heading": "1. Introduction", "paragraph_index": 3, "text": "..."}, ...]
    """
    if not os.path.exists(file_path):
        raise ExtractionError(f"File not found: {file_path}")

    try:
        doc = DocxDocument(file_path)
    except Exception as e:
        raise ExtractionError(f"Failed to parse DOCX '{file_path}': {str(e)}")

    if len(doc.paragraphs) == 0:
        raise ExtractionError(f"DOCX file '{file_path}' is empty.")

    paragraphs = []
    current_heading = "Preamble"

    for idx, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if not text:
            continue

        style_name = (para.style.name or "").lower()
        is_heading = style_name.startswith("heading") or style_name == "title"

        if is_heading:
            current_heading = text
            continue  # heading itself isn't stored as content, just tracked

        paragraphs.append({
            "heading": current_heading,
            "paragraph_index": idx,
            "text": text
        })

    # Tables (requirement lists, escalation matrices ...) used to be skipped
    # entirely. Each row now becomes its own traceable block.
    for t_idx, table in enumerate(doc.tables, start=1):
        for r_idx, row in enumerate(table.rows, start=1):
            cells = []
            for cell in row.cells:
                value = cell.text.strip()
                if value and value not in cells:
                    cells.append(value)
            if cells:
                paragraphs.append({
                    "heading": f"Table {t_idx}",
                    "paragraph_index": f"table {t_idx}, row {r_idx}",
                    "text": " | ".join(cells),
                })

    if not paragraphs:
        raise ExtractionError(
            f"No usable paragraph content found in '{file_path}' "
            f"(document may contain only headings/images)."
        )

    return paragraphs


def find_hidden_docx_text(file_path: str) -> list[dict]:
    """
    Finds text a human reader cannot see but a model would read: runs
    formatted as hidden, or white / near-white font colour, or tiny fonts.
    Used by the prompt-injection guard (SRS Step 43: "hidden conflicting
    instructions").
    """
    try:
        doc = DocxDocument(file_path)
    except Exception:
        return []
    found = []
    for idx, para in enumerate(doc.paragraphs):
        for run in para.runs:
            text = (run.text or "").strip()
            if not text:
                continue
            font = run.font
            hidden = bool(font.hidden)
            color = None
            try:
                color = str(font.color.rgb) if font.color is not None and font.color.rgb is not None else None
            except Exception:
                color = None
            white = color is not None and color.upper() in {"FFFFFF", "FEFEFE", "FDFDFD", "FAFAFA"}
            tiny = font.size is not None and font.size.pt is not None and font.size.pt < 2
            if hidden or white or tiny:
                found.append({
                    "text": text,
                    "location": f"paragraph {idx}",
                    "reason": "hidden" if hidden else ("white text" if white else "tiny font"),
                })
    return found


def extract_text_file(file_path: str) -> list[dict]:
    """TXT / Markdown support (SRS Step 4 optional formats)."""
    if not os.path.exists(file_path):
        raise ExtractionError(f"File not found: {file_path}")
    try:
        with open(file_path, encoding="utf-8", errors="replace") as f:
            raw = f.read()
    except Exception as e:
        raise ExtractionError(f"Failed to read '{file_path}': {str(e)}")

    blocks, heading, buffer, line_no = [], "Preamble", [], 0
    start_line = 1

    def flush():
        text = " ".join(buffer).strip()
        if text:
            blocks.append({"heading": heading, "paragraph_index": f"line {start_line}", "text": text})

    for line_no, line in enumerate(raw.splitlines(), start=1):
        stripped = line.strip()
        if stripped.startswith("#"):
            flush(); buffer = []
            heading = stripped.lstrip("#").strip() or heading
            start_line = line_no + 1
        elif not stripped:
            flush(); buffer = []
            start_line = line_no + 1
        else:
            if not buffer:
                start_line = line_no
            buffer.append(stripped)
    flush()

    if not blocks:
        raise ExtractionError(f"'{file_path}' is empty.")
    return blocks


def extract_document(file_path: str) -> dict:
    """
    Main entry point. Detects file type by extension and routes to the
    correct extractor. Returns a normalized structure:

    {
        "file_type": "pdf" | "docx",
        "file_path": "...",
        "content": [ ... ]   # list of page dicts or paragraph dicts
    }
    """
    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".pdf":
        content = extract_pdf(file_path)
        return {"file_type": "pdf", "file_path": file_path, "content": content}

    elif ext == ".docx":
        content = extract_docx(file_path)
        return {
            "file_type": "docx", "file_path": file_path, "content": content,
            "hidden_text": find_hidden_docx_text(file_path),
        }

    elif ext in (".txt", ".md"):
        content = extract_text_file(file_path)
        # same block shape as DOCX, so the chunker treats it identically
        return {"file_type": "docx", "file_path": file_path, "content": content}

    else:
        raise ExtractionError(
            f"Unsupported file type '{ext}'. Supported: .pdf, .docx, .txt, .md."
        )


if __name__ == "__main__":
    # Quick manual test — run: python extractor.py path/to/file.docx
    import sys
    if len(sys.argv) < 2:
        print("Usage: python extractor.py <file_path>")
        sys.exit(1)

    result = extract_document(sys.argv[1])
    print(f"File type: {result['file_type']}")
    print(f"Extracted {len(result['content'])} content blocks:\n")
    for block in result["content"][:5]:
        print(block)
        print("---")